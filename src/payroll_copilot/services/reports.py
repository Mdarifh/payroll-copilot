"""Payroll reports from stored payslips and records (never recalculated): one table per
report, exported as JSON (screen), CSV, Excel or PDF.

Every report is (columns, rows, totals) for a month; money stays Decimal until export.
"""

import csv
import io
from collections.abc import Callable
from dataclasses import dataclass, field
from decimal import Decimal
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from payroll_copilot.core.dates import PayrollPeriod
from payroll_copilot.core.exceptions import NotFoundError, PayrollCopilotError
from payroll_copilot.db.models import (
    Employee,
    LopRecord,
    OvertimeEntry,
    PayAdjustment,
    PayrollRun,
)
from payroll_copilot.db.models.adjustments import ACTIVE
from payroll_copilot.repositories.attendance_repo import AttendanceRepository

ZERO = Decimal(0)


@dataclass
class Report:
    key: str
    title: str
    period: str
    columns: list[tuple[str, str, str]]  # (key, label, kind: text | money | days | number)
    rows: list[dict[str, Any]]
    totals: dict[str, Any] = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)

    def as_dict(self) -> dict[str, Any]:
        def out(v: Any) -> Any:
            return str(v) if isinstance(v, Decimal) else v

        return {"key": self.key, "title": self.title, "period": self.period,
                "columns": [{"key": k, "label": label, "kind": kind}
                            for k, label, kind in self.columns],
                "rows": [{k: out(v) for k, v in r.items()} for r in self.rows],
                "totals": {k: out(v) for k, v in self.totals.items()},
                "notes": self.notes}  # fmt: skip


def _runs(session: Session, period: PayrollPeriod, department: str | None):
    stmt = (select(PayrollRun, Employee).join(Employee, Employee.id == PayrollRun.employee_id)
            .where(PayrollRun.period == str(period)).order_by(Employee.employee_code))
    if department:
        stmt = stmt.where(Employee.department == department)
    return session.execute(stmt).all()


def _lines(slip: dict[str, Any], bucket: str) -> dict[str, Decimal]:
    out: dict[str, Decimal] = {}
    for ln in slip.get(bucket, []):
        out[ln["head_code"]] = out.get(ln["head_code"], ZERO) + Decimal(ln["amount"])
    return out


def _sum(rows: list[dict[str, Any]], keys: list[str]) -> dict[str, Any]:
    return {k: sum((r[k] for r in rows if isinstance(r.get(k), Decimal)), ZERO) for k in keys}


EMP_COLS = [("employee_code", "Code", "text"), ("employee_name", "Employee", "text"),
            ("department", "Department", "text")]  # fmt: skip


def _base(e: Employee) -> dict[str, Any]:
    return {"employee_code": e.employee_code, "employee_name": e.full_name,
            "department": e.department or ""}


def payroll_register(session, period, department=None) -> Report:
    rows = []
    for run, e in _runs(session, period, department):
        d = _lines(run.payslip, "deductions")
        rows.append({**_base(e), "paid_days": run.paid_days, "lop_days": run.lop_days,
                     "gross": run.gross_earnings, "pf": d.get("EE_PF", ZERO),
                     "esic": d.get("EE_ESIC", ZERO), "pt": d.get("PT", ZERO),
                     "tds": d.get("TDS", ZERO),
                     "other": run.total_deductions - sum(
                         (d.get(h, ZERO) for h in ("EE_PF", "EE_ESIC", "PT", "TDS")), ZERO),
                     "deductions": run.total_deductions, "net": run.net_pay,
                     "status": run.status.value})  # fmt: skip
    money = ["gross", "pf", "esic", "pt", "tds", "other", "deductions", "net"]
    return Report("payroll_register", "Payroll register", str(period), [
        *EMP_COLS, ("paid_days", "Paid days", "days"), ("lop_days", "LOP", "days"),
        ("gross", "Gross", "money"), ("pf", "PF", "money"), ("esic", "ESIC", "money"),
        ("pt", "PT", "money"), ("tds", "TDS", "money"), ("other", "Other ded.", "money"),
        ("deductions", "Total ded.", "money"), ("net", "Net pay", "money"),
        ("status", "Status", "text")], rows, _sum(rows, money))  # fmt: skip


def pf_report(session, period, department=None) -> Report:
    rows = []
    for run, e in _runs(session, period, department):
        d, er = _lines(run.payslip, "deductions"), _lines(run.payslip, "employer_contributions")
        if not (d.get("EE_PF") or er.get("ER_EPF")):
            continue
        rows.append({**_base(e), "uan": e.uan or "", "ee_pf": d.get("EE_PF", ZERO),
                     "er_epf": er.get("ER_EPF", ZERO), "eps": er.get("EPS", ZERO),
                     "edli": er.get("EDLI", ZERO), "admin": er.get("EPF_ADMIN", ZERO)})
    money = ["ee_pf", "er_epf", "eps", "edli", "admin"]
    return Report("pf", "PF report (ECR)", str(period), [
        *EMP_COLS, ("uan", "UAN", "text"), ("ee_pf", "Employee PF", "money"),
        ("er_epf", "Employer EPF", "money"), ("eps", "EPS", "money"), ("edli", "EDLI", "money"),
        ("admin", "EPF admin", "money")], rows, _sum(rows, money),
        ["Establishment-level admin minimum is applied in the EPF challan (Compliance)."])


def esic_report(session, period, department=None) -> Report:
    rows = []
    for run, e in _runs(session, period, department):
        d, er = _lines(run.payslip, "deductions"), _lines(run.payslip, "employer_contributions")
        if not (d.get("EE_ESIC") or er.get("ER_ESIC")):
            continue
        rows.append({**_base(e), "ip": e.esic_ip_number or "", "gross": run.gross_earnings,
                     "ee": d.get("EE_ESIC", ZERO), "er": er.get("ER_ESIC", ZERO)})
    return Report("esic", "ESIC report", str(period), [
        *EMP_COLS, ("ip", "IP number", "text"), ("gross", "Gross", "money"),
        ("ee", "Employee ESIC", "money"), ("er", "Employer ESIC", "money")], rows,
        _sum(rows, ["gross", "ee", "er"]))  # fmt: skip


def pt_report(session, period, department=None) -> Report:
    rows = []
    for run, e in _runs(session, period, department):
        pt = _lines(run.payslip, "deductions").get("PT", ZERO)
        rows.append({**_base(e), "state": e.work_state, "gross": run.gross_earnings, "pt": pt})
    return Report("pt", "Professional Tax report", str(period), [
        *EMP_COLS, ("state", "State", "text"), ("gross", "Gross", "money"),
        ("pt", "PT", "money")], rows, _sum(rows, ["gross", "pt"]))  # fmt: skip


def tds_report(session, period, department=None) -> Report:
    rows = []
    for run, e in _runs(session, period, department):
        rows.append({**_base(e), "pan": e.pan or "", "regime": e.tax_regime.value,
                     "gross": run.gross_earnings,
                     "tds": _lines(run.payslip, "deductions").get("TDS", ZERO)})
    return Report("tds", "TDS report", str(period), [
        *EMP_COLS, ("pan", "PAN", "text"), ("regime", "Regime", "text"),
        ("gross", "Gross", "money"), ("tds", "TDS", "money")], rows,
        _sum(rows, ["gross", "tds"]))  # fmt: skip


def lop_report(session, period, department=None) -> Report:
    repo = AttendanceRepository(session)
    stmt = select(Employee).order_by(Employee.employee_code)
    if department:
        stmt = stmt.where(Employee.department == department)
    records = {r.employee_id: r for r in session.scalars(select(LopRecord).where(
        LopRecord.period == str(period)))}
    rows = []
    for e in session.scalars(stmt):
        absent = sum((x.days for x in repo.exceptions(e.id, period)), ZERO)
        rec = records.get(e.id)
        if not absent and rec is None:
            continue
        by = session.get(Employee, rec.updated_by_id) if rec else None
        rows.append({**_base(e), "absent": absent, "lop": rec.days if rec else ZERO,
                     "decided": "yes" if rec else "NO", "reason": rec.reason if rec else "",
                     "by": by.full_name if by else ""})  # fmt: skip
    return Report("lop", "LOP report", str(period), [
        *EMP_COLS, ("absent", "Absent (no leave)", "days"), ("lop", "LOP days", "days"),
        ("decided", "Decided", "text"), ("reason", "Reason", "text"), ("by", "By", "text")],
        rows, _sum(rows, ["absent", "lop"]))  # fmt: skip


def overtime_report(session, period, department=None) -> Report:
    stmt = (select(OvertimeEntry, Employee).join(Employee, Employee.id == OvertimeEntry.employee_id)
            .where(OvertimeEntry.period == str(period), OvertimeEntry.status == ACTIVE)
            .order_by(Employee.employee_code, OvertimeEntry.work_date))  # fmt: skip
    if department:
        stmt = stmt.where(Employee.department == department)
    rows = [{**_base(e), "date": o.work_date.isoformat(), "hours": o.hours,
             "rate": o.rate_per_hour, "amount": o.amount, "source": o.rate_source}
            for o, e in session.execute(stmt)]  # fmt: skip
    return Report("overtime", "Overtime report", str(period), [
        *EMP_COLS, ("date", "Date", "text"), ("hours", "Hours", "number"),
        ("rate", "Rate / hour", "money"), ("amount", "Amount", "money"),
        ("source", "Rate source", "text")], rows, _sum(rows, ["hours", "amount"]))


def one_off_report(session, period, department=None) -> Report:
    stmt = (select(PayAdjustment, Employee).join(Employee, Employee.id == PayAdjustment.employee_id)
            .where(PayAdjustment.period == str(period), PayAdjustment.status == ACTIVE)
            .order_by(Employee.employee_code))  # fmt: skip
    if department:
        stmt = stmt.where(Employee.department == department)
    rows = [{**_base(e), "kind": a.kind.title(), "head": a.label or a.head_code,
             "amount": a.amount if a.kind == "EARNING" else -a.amount, "reason": a.reason}
            for a, e in session.execute(stmt)]  # fmt: skip
    return Report("one_off", "Bonus / miscellaneous pay report", str(period), [
        *EMP_COLS, ("kind", "Type", "text"), ("head", "Component", "text"),
        ("amount", "Amount", "money"), ("reason", "Reason", "text")], rows,
        _sum(rows, ["amount"]))  # fmt: skip


def arrears_report(session, period, department=None) -> Report:
    rows = []
    for run, e in _runs(session, period, department):
        for ln in run.payslip.get("earnings", []):
            if ln.get("is_arrear"):
                rows.append({**_base(e), "head": ln["name"], "amount": Decimal(ln["amount"])})
    return Report("arrears", "Arrears report", str(period), [
        *EMP_COLS, ("head", "Component", "text"), ("amount", "Arrear", "money")], rows,
        _sum(rows, ["amount"]), ["Negative amounts are recoveries."])


def variance_report(session, period, department=None, threshold: Decimal = Decimal(10)) -> Report:
    prev = {r.employee_id: r for r in session.scalars(select(PayrollRun).where(
        PayrollRun.period == str(period.previous())))}
    rows = []
    for run, e in _runs(session, period, department):
        old = prev.get(e.id)
        if old is None or old.net_pay == 0:
            continue
        change = (run.net_pay - old.net_pay) * 100 / old.net_pay
        if abs(change) < threshold:
            continue
        rows.append({**_base(e), "previous": old.net_pay, "current": run.net_pay,
                     "difference": run.net_pay - old.net_pay,
                     "percent": change.quantize(Decimal("0.1")),
                     "lop": f"{old.lop_days} → {run.lop_days}",
                     "gross": f"{old.gross_earnings} → {run.gross_earnings}"})  # fmt: skip
    return Report("variance", f"Salary variance (net change ≥ {threshold}%)", str(period), [
        *EMP_COLS, ("previous", f"Net {period.previous().label}", "money"),
        ("current", f"Net {period.label}", "money"), ("difference", "Difference", "money"),
        ("percent", "%", "number"), ("lop", "LOP", "text"), ("gross", "Gross", "text")],
        rows, _sum(rows, ["previous", "current", "difference"]))  # fmt: skip


def anomaly_report(session, period, department=None) -> Report:
    from payroll_copilot.services import payroll_audit as pa

    cols = [("employee_code", "Code", "text"), ("employee_name", "Employee", "text"),
            ("issue", "Issue", "text"), ("severity", "Severity", "text"),
            ("detail", "Root cause / detail", "text")]  # fmt: skip
    run = pa.latest(session, period)
    if run is None:
        return Report("anomaly", "Anomaly report (AI audit)", str(period), cols, [], {},
                      [f"No AI audit for {period.label} yet - run one from AI Audit."])
    rows = [{"employee_code": i["employee_code"], "employee_name": i["employee_name"],
             "issue": i["issue"], "severity": i["severity"],
             "detail": (i.get("causes") or [i.get("message") or i.get("title") or ""])[0]}
            for i in [*(run.findings or []), *(run.anomalies or [])]]  # fmt: skip
    return Report("anomaly", "Anomaly report (AI audit)", str(period), cols, rows, {},
        [f"From the AI audit finished {run.finished_at:%d %b %Y %H:%M} UTC."])


REPORTS: dict[str, tuple[str, Callable[..., Report]]] = {
    "payroll_register": ("Payroll register", payroll_register),
    "pf": ("PF report (ECR)", pf_report), "esic": ("ESIC report", esic_report),
    "pt": ("Professional Tax report", pt_report), "tds": ("TDS report", tds_report),
    "lop": ("LOP report", lop_report), "overtime": ("Overtime report", overtime_report),
    "one_off": ("Bonus / miscellaneous pay", one_off_report),
    "arrears": ("Arrears report", arrears_report),
    "variance": ("Salary variance", variance_report),
    "anomaly": ("Anomaly report (AI audit)", anomaly_report),
}  # fmt: skip


def build(session: Session, key: str, period: PayrollPeriod,
          department: str | None = None) -> Report:  # fmt: skip
    if key not in REPORTS:
        raise NotFoundError(f"Unknown report {key}")
    return REPORTS[key][1](session, period, department)


# ---------------------------------------------------------------------- export


def _cell(value: Any, kind: str) -> Any:
    if isinstance(value, Decimal):
        return float(value) if kind in ("money", "number", "days") else str(value)
    return value


def to_csv(report: Report) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow([label for _, label, _ in report.columns])
    for r in report.rows:
        writer.writerow([str(r.get(k, "")) for k, _, _ in report.columns])
    if report.totals:
        writer.writerow([("Total" if i == 0 else str(report.totals.get(k, "")))
                         for i, (k, _, _) in enumerate(report.columns)])
    return buffer.getvalue().encode("utf-8-sig")  # BOM: Excel opens ₹ / Hindi names right


def to_xlsx(report: Report) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = report.key[:31]
    ws.append([f"{report.title} · {PayrollPeriod.parse(report.period).label}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append([label for _, label, _ in report.columns])
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
    for r in report.rows:
        ws.append([_cell(r.get(k), kind) for k, _, kind in report.columns])
    if report.totals:
        ws.append([("Total" if i == 0 else _cell(report.totals.get(k), kind))
                   for i, (k, _, kind) in enumerate(report.columns)])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    for i, (_, _, kind) in enumerate(report.columns, 1):
        if kind == "money":
            for row in ws.iter_rows(min_row=4, min_col=i, max_col=i):
                row[0].number_format = "#,##,##0.00"
    for column in ws.columns:
        width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
        ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 8), 40)
    for note in report.notes:
        ws.append([])
        ws.append([note])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(report: Report) -> bytes:
    from fpdf import FPDF

    def latin(text: str) -> str:  # built-in PDF fonts: Latin-1 only
        for a, b in (("₹", "Rs."), ("→", "->"), ("≥", ">="), ("≤", "<="), ("–", "-"),
                     ("—", "-"), ("·", "-"), ("“", '"'), ("”", '"'), ("’", "'")):
            text = text.replace(a, b)
        return text.encode("latin-1", "replace").decode("latin-1")

    cols = report.columns
    pdf = FPDF(orientation="L" if len(cols) > 6 else "P", format="A4")
    family = "Helvetica"
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    pdf.set_font(family, size=13)
    pdf.cell(0, 8, latin(f"{report.title} - {PayrollPeriod.parse(report.period).label}"),
             new_x="LMARGIN", new_y="NEXT")  # fmt: skip
    usable = pdf.w - pdf.l_margin - pdf.r_margin
    width = usable / len(cols)
    pdf.set_font(family, size=7)
    pdf.set_fill_color(79, 70, 229)
    pdf.set_text_color(255, 255, 255)
    for _, label, _ in cols:
        pdf.cell(width, 6, latin(label[:22]), border=1, fill=True)
    pdf.ln()
    pdf.set_text_color(0, 0, 0)

    def line(values: list[str]) -> None:
        for (_, _, kind), v in zip(cols, values, strict=True):
            pdf.cell(width, 5, latin(v[:28]), border=1, align="R" if kind != "text" else "L")
        pdf.ln()

    for r in report.rows:
        line([str(r.get(k, "")) for k, _, _ in cols])
    if report.totals:
        line([("Total" if i == 0 else str(report.totals.get(k, ""))) for i, (k, _, _) in
              enumerate(cols)])  # fmt: skip
    for note in report.notes:
        pdf.ln(2)
        pdf.multi_cell(0, 5, latin(note))
    return bytes(pdf.output())


EXPORTS = {"csv": ("text/csv", to_csv), "xlsx": (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", to_xlsx),
    "pdf": ("application/pdf", to_pdf)}


def export(report: Report, fmt: str) -> tuple[bytes, str, str]:
    if fmt not in EXPORTS:
        raise PayrollCopilotError("Format must be csv, xlsx or pdf")
    media, fn = EXPORTS[fmt]
    return fn(report), media, f"{report.key}_{report.period}.{fmt}"
